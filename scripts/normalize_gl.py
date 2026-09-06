"""Dataset 02 adapter — investor-level GL → loader, as the same canonical close.

Same contract as Dataset 01 (D-011 / D-012):
  raw GL facts → evaluate_rules(crosswalks) → compare(loader) → origin

  Source   = GL extract rows + loader workbook sheets
  Line     = one GL movement group (Legal Entity × GL Account × Trans Type × entity currency)
  Rule     = one crosswalk row (CoA Mapping / LE Mapping / Deal Mapping) — exact hit only
  Case     = movement group with no exact crosswalk hit, unmapped entity / deal,
             or a crosswalk result the verified loader disagrees with
  TieOut   = GL per-batch zero balance, loader per-batch zero balance,
             GL vs loader entity net, README residue counts

The 34k GL rows are read directly with openpyxl (read-only) and grouped; the
full body is never written under data/extracted. Nothing under data/raw is modified.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Optional

from canonical_model import (
    CanonicalClose,
    Case,
    Close,
    Line,
    Rule,
    Source,
    TieOut,
    ensure_close_tasks,
    primary_reason_of,
    priority_for_reason,
)
from extract_excel import make_source_id

CLOSE_ID = "CLOSE-2026-Q2-KESTREL-GL-LOADER"
FUND_ID = "MULTI-ENTITY-KESTREL"
PERIOD = "2026-Q2"
WORKFLOW = "gl_to_loader"
CONTRACT_VERSION = "v1"

GL_SHEET = "Investor-Level GL"
UPLOAD_SHEET = "Upload Template (VERIFIED v4c)"
COA_MAP_SHEET = "CoA Mapping"
LE_MAP_SHEET = "LE Mapping"
DEAL_MAP_SHEET = "Deal Mapping"
INVESTOR_MAP_SHEET = "Investor Mapping"
ENTITY_LIST_SHEET = "Entity Listing"
DEALS_LIST_SHEET = "Deals List"
INVESTORS_LIST_SHEET = "Investors List"
MAPPING_GAPS_SHEET = "Mapping Gaps"
MOVEMENTS_SHEET = "Movements Rec"
RULE_VERSION = "v4c"  # verified loader workbook version

# Immutable Dataset 02 README residue (README2.md "Known unmatched rows").
EXPECTED_RESIDUE = {
    "upload_entities_not_in_entity_listing": 4,
    "upload_deals_not_in_deals_list": 16,
    "investor_mapping_names_not_in_investors_list": 198,
}

TOL = Decimal("0.01")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------


def _s(v: Any) -> Optional[str]:
    if v is None:
        return None
    text = str(v).strip()
    return text or None


def _d(v: Any) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    try:
        return Decimal(str(v))
    except Exception:  # noqa: BLE001
        return Decimal("0")


def _money(v: Decimal) -> float:
    return float(v.quantize(Decimal("0.01")))


@dataclass
class _Group:
    legal_entity: str
    gl_account: str
    trans_type: str
    currency: str
    account_type: Optional[str] = None
    row_count: int = 0
    debits: Decimal = Decimal("0")
    credits: Decimal = Decimal("0")
    net: Decimal = Decimal("0")
    first_row: Optional[int] = None
    last_row: Optional[int] = None
    sample_rows: list[int] = field(default_factory=list)
    deals: set[str] = field(default_factory=set)
    investors: set[str] = field(default_factory=set)
    batch_types: set[str] = field(default_factory=set)
    batch_ids: set[str] = field(default_factory=set)


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def _find_file(input_dir: Path, needle: str) -> Path:
    hits = [p for p in input_dir.rglob("*.xlsx") if needle.lower() in p.name.lower() and not p.name.startswith("~$")]
    if not hits:
        raise SystemExit(f"Dataset 02: no workbook matching {needle!r} under {input_dir}")
    return sorted(hits)[0]


def _sheet(wb, name: str, skip: int = 0) -> tuple[list[Any], list[tuple[int, tuple]]]:
    """Return (header, [(row_number, values)]) skipping fully empty rows."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[skip]) if rows else []
    body = [
        (i + 1, r)
        for i, r in enumerate(rows)
        if i > skip and any(c is not None and c != "" for c in r)
    ]
    return header, body


def read_gl_groups(gl_path: Path) -> tuple[dict[tuple, _Group], dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(gl_path, data_only=True, read_only=True)
    ws = wb[GL_SHEET]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    idx: dict[str, int] = {}
    for i, h in enumerate(header):
        if h is not None and h not in idx:
            idx[str(h)] = i

    def col(row, name):
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    groups: dict[tuple, _Group] = {}
    batches: dict[tuple, list[Decimal]] = defaultdict(lambda: [Decimal("0"), Decimal("0")])
    ccy_totals: dict[str, list[Decimal]] = defaultdict(lambda: [Decimal("0"), Decimal("0")])
    row_total = 0
    for row_number, row in enumerate(it, start=2):
        le = _s(col(row, "Legal Entity"))
        if not le:
            continue
        row_total += 1
        acct = _s(col(row, "GL Account")) or ""
        tt = _s(col(row, "Trans Type")) or ""
        ccy = _s(col(row, "Legal Entity Currency")) or "XXX"
        key = (le, acct, tt, ccy)
        g = groups.get(key)
        if g is None:
            g = groups[key] = _Group(le, acct, tt, ccy, account_type=_s(col(row, "Account Type")))
            g.first_row = row_number
        deb = _d(col(row, "Debits (Entity Currency)"))
        cre = _d(col(row, "Credits (Entity Currency)"))
        g.row_count += 1
        g.debits += deb
        g.credits += cre
        g.net += _d(col(row, "Amount (Entity Currency)"))
        g.last_row = row_number
        if len(g.sample_rows) < 5:
            g.sample_rows.append(row_number)
        deal = _s(col(row, "Deal Name"))
        if deal:
            g.deals.add(deal)
        inv = _s(col(row, "Investor"))
        if inv:
            g.investors.add(inv)
        bt = _s(col(row, "Batch Type"))
        if bt:
            g.batch_types.add(bt)
        bid = _s(col(row, "Batch ID"))
        if bid:
            g.batch_ids.add(bid)
        b = batches[(le, bid)]
        b[0] += deb
        b[1] += cre
        ccy_totals[ccy][0] += deb
        ccy_totals[ccy][1] += cre
    wb.close()
    meta = {
        "row_count": row_total,
        "batches": {f"{k[0]}|{k[1]}": [str(v[0]), str(v[1])] for k, v in batches.items()},
        "ccy_totals": {k: [str(v[0]), str(v[1])] for k, v in ccy_totals.items()},
    }
    return groups, meta


# ---------------------------------------------------------------------------
# Normalize
# ---------------------------------------------------------------------------


def normalize_gl_to_canonical_model(input_dir: Path) -> CanonicalClose:
    from openpyxl import load_workbook

    gl_path = _find_file(input_dir, "Investor-Level GL")
    loader_path = _find_file(input_dir, "verified loader")
    sample_path = None
    try:
        sample_path = _find_file(input_dir, "Phase I loader")
    except SystemExit:
        pass

    groups, gl_meta = read_gl_groups(gl_path)
    wb = load_workbook(loader_path, data_only=True, read_only=True)

    # ---- Sources -----------------------------------------------------------
    sources: list[Source] = []
    gl_source_id = make_source_id(gl_path, GL_SHEET)
    sources.append(
        Source(gl_source_id, gl_path.name, "xlsx", "gl_extract", {"sheet": GL_SHEET, "rows": gl_meta["row_count"]})
    )
    roles = {
        UPLOAD_SHEET: "loader_output",
        COA_MAP_SHEET: "loader_mapping",
        LE_MAP_SHEET: "loader_mapping",
        DEAL_MAP_SHEET: "loader_mapping",
        INVESTOR_MAP_SHEET: "loader_mapping",
        ENTITY_LIST_SHEET: "reference_list",
        DEALS_LIST_SHEET: "reference_list",
        INVESTORS_LIST_SHEET: "reference_list",
        MAPPING_GAPS_SHEET: "working_file",
        MOVEMENTS_SHEET: "working_file",
    }
    sid: dict[str, str] = {}
    for sheet_name in wb.sheetnames:
        s_id = make_source_id(loader_path, sheet_name)
        sid[sheet_name] = s_id
        sources.append(
            Source(s_id, loader_path.name, "xlsx", roles.get(sheet_name, "working_file"), {"sheet": sheet_name})
        )
    if sample_path is not None:
        sources.append(Source(make_source_id(sample_path, "Sheet1"), sample_path.name, "xlsx", "loader_sample", {"sheet": "Sheet1"}))

    # ---- Rules (crosswalks, exact hit only) --------------------------------
    rules: list[Rule] = []
    coa_by_key: dict[tuple, Rule] = {}
    coa_by_tt: dict[str, Rule] = {}
    _, coa_rows = _sheet(wb, COA_MAP_SHEET)
    for row_number, r in coa_rows:
        acct, tt = _s(r[0]), _s(r[2])
        if not tt:
            continue
        target_code = _s(r[3])
        target_tt = _s(r[6])
        if not target_tt and not target_code:
            continue
        rule = Rule(
            rule_id=f"RULE-COA-{len(coa_by_key) + len(coa_by_tt) + 1:03d}",
            name=f"CoA Mapping: {acct or '*'} / {tt}",
            condition={
                "field": "gl_account,trans_type",
                "operator": "equals",
                "value": [acct, tt],
                "source_id": sid[COA_MAP_SHEET],
                "row": row_number,
            },
            action={
                k: v
                for k, v in {
                    "account": target_code,
                    "account_name": _s(r[4]),
                    "classification": target_tt,
                    "trans_type_debit": _s(r[7]),
                    "batch_type": _s(r[12]),
                }.items()
                if v is not None
            },
            origin="coa_mapping",
            status="active",
            version=RULE_VERSION,
            kind="mapping",
        )
        rules.append(rule)
        if acct:
            coa_by_key[(acct, tt)] = rule
        else:
            coa_by_tt[tt] = rule

    le_rules: dict[str, Rule] = {}
    _, le_rows = _sheet(wb, LE_MAP_SHEET, skip=1)
    for row_number, r in le_rows:
        le = _s(r[1])
        if not le or not _s(r[4]):
            continue
        rule = Rule(
            rule_id=f"RULE-LE-{len(le_rules) + 1:03d}",
            name=f"LE Mapping: {le}",
            condition={"field": "legal_entity", "operator": "equals", "value": le, "source_id": sid[LE_MAP_SHEET], "row": row_number},
            action={"fund": _s(r[4]), "fund_name": _s(r[3]), "fund_currency": _s(r[5])},
            origin="le_mapping",
            status="active",
            version=RULE_VERSION,
        )
        rules.append(rule)
        le_rules[le] = rule

    deal_rules: dict[str, Rule] = {}
    _, deal_rows = _sheet(wb, DEAL_MAP_SHEET)
    for row_number, r in deal_rows:
        deal = _s(r[0])
        if not deal or deal in deal_rules or not _s(r[6]):
            continue
        rule = Rule(
            rule_id=f"RULE-DEAL-{len(deal_rules) + 1:03d}",
            name=f"Deal Mapping: {deal}",
            condition={"field": "deal_name", "operator": "equals", "value": deal, "source_id": sid[DEAL_MAP_SHEET], "row": row_number},
            action={"deal_id": _s(r[6]), "deal_name": _s(r[5])},
            origin="deal_mapping",
            status="active",
            version=RULE_VERSION,
        )
        rules.append(rule)
        deal_rules[deal] = rule

    # ---- Reference lists + workbook residue --------------------------------
    _, ent_rows = _sheet(wb, ENTITY_LIST_SHEET)
    entity_listing = {_s(r[1]) for r in (x[1] for x in ent_rows) if _s(r[1])}
    _, dl_rows = _sheet(wb, DEALS_LIST_SHEET)
    deals_list = {_s(r[0]) for r in (x[1] for x in dl_rows) if _s(r[0])}
    _, il_rows = _sheet(wb, INVESTORS_LIST_SHEET)
    investors_list = {_s(r[1]) for r in (x[1] for x in il_rows) if _s(r[1])}
    _, im_rows = _sheet(wb, INVESTOR_MAP_SHEET)
    investor_map_names = [_s(r[4]) for _, r in im_rows if _s(r[4])]
    unmatched_investors = sorted({n for n in investor_map_names if n not in investors_list})

    gaps: dict[tuple, dict[str, Any]] = {}
    _, gap_rows = _sheet(wb, MAPPING_GAPS_SHEET)
    for row_number, r in gap_rows:
        acct, tt = _s(r[0]), _s(r[1])
        if acct and tt:
            gaps[(acct, tt)] = {
                "row": row_number,
                "row_count": r[2],
                "total": r[3],
                "proposed_account": _s(r[4]),
                "proposed_trans_type": _s(r[5]),
                "approval": _s(r[6]),
            }

    # ---- Loader (filed snapshot) --------------------------------------------
    up_hdr, up_rows = _sheet(wb, UPLOAD_SHEET)
    up_idx = {str(h): i for i, h in enumerate(up_hdr) if h is not None}
    upload_by_le_tt: dict[tuple, list[Decimal]] = defaultdict(lambda: [Decimal("0"), Decimal("0")])
    upload_by_le_net: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    upload_batches: dict[Any, list[Decimal]] = defaultdict(lambda: [Decimal("0"), Decimal("0")])
    upload_entities: set[str] = set()
    upload_deals: set[str] = set()
    for _, r in up_rows:
        le = _s(r[up_idx["Legal Entity"]])
        if not le:
            continue
        upload_entities.add(le)
        deal = _s(r[up_idx["Deal Name"]])
        if deal:
            upload_deals.add(deal)
        tt = _s(r[up_idx["Trans Type"]]) or ""
        amt = _d(r[up_idx["Investor Amount (LE)"]])
        is_debit = str(r[up_idx["Is Debit"]]).strip().upper() == "Y"
        agg = upload_by_le_tt[(le, tt)]
        b = upload_batches[r[up_idx["Batch Index"]]]
        if is_debit:
            agg[0] += amt
            b[0] += amt
            upload_by_le_net[le] += amt
        else:
            agg[1] += amt
            b[1] += amt
            upload_by_le_net[le] -= amt
    wb.close()

    # ---- Evaluate rules per movement group ---------------------------------
    ordered = sorted(groups.values(), key=lambda g: (g.legal_entity, g.gl_account, g.trans_type, g.currency))

    # Crosswalk result aggregated at (LE, target trans type) — the grain the
    # loader can be compared at. Each Line inherits the agreement of its key.
    gl_by_le_tt: dict[tuple, list[Decimal]] = defaultdict(lambda: [Decimal("0"), Decimal("0")])
    hits: dict[tuple, Optional[Rule]] = {}
    for g in ordered:
        rule = coa_by_key.get((g.gl_account, g.trans_type)) or coa_by_tt.get(g.trans_type)
        hits[(g.legal_entity, g.gl_account, g.trans_type, g.currency)] = rule
        if rule is None:
            continue
        t_default = rule.action.get("classification")
        t_debit = rule.action.get("trans_type_debit") or t_default
        gl_by_le_tt[(g.legal_entity, t_debit)][0] += g.debits
        gl_by_le_tt[(g.legal_entity, t_default)][1] += g.credits

    agreement: dict[tuple, Optional[bool]] = {}
    for key, (deb, cre) in gl_by_le_tt.items():
        le = key[0]
        if le not in upload_entities:
            agreement[key] = None  # entity not in this loader tranche — nothing filed
            continue
        u = upload_by_le_tt.get(key)
        if u is None:
            agreement[key] = False
            continue
        agreement[key] = abs(u[0] - deb) <= TOL and abs(u[1] - cre) <= TOL

    lines: list[Line] = []
    cases: list[Case] = []
    gl_net_by_le: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    line_ccy_totals: dict[str, list[Decimal]] = defaultdict(lambda: [Decimal("0"), Decimal("0")])
    evaluation_report: list[dict[str, Any]] = []

    for n, g in enumerate(ordered, start=1):
        line_id = f"GL-{n:04d}"
        key = (g.legal_entity, g.gl_account, g.trans_type, g.currency)
        coa_rule = hits[key]
        le_rule = le_rules.get(g.legal_entity)
        deal_hits = {d: deal_rules.get(d) for d in sorted(g.deals)}
        missing_deals = [d for d, r in deal_hits.items() if r is None]
        gl_net_by_le[g.legal_entity] += g.net
        line_ccy_totals[g.currency][0] += g.debits
        line_ccy_totals[g.currency][1] += g.credits

        raw_facts = {
            "legal_entity": g.legal_entity,
            "gl_account": g.gl_account,
            "account_type": g.account_type,
            "trans_type": g.trans_type,
            "currency": g.currency,
            "row_count": g.row_count,
            "debits": _money(g.debits),
            "credits": _money(g.credits),
            "net": _money(g.net),
            "batch_types": sorted(g.batch_types),
            "batch_count": len(g.batch_ids),
            "deal_names": sorted(g.deals),
            "investor_count": len(g.investors),
            "gl_rows": {"first": g.first_row, "last": g.last_row, "sample": g.sample_rows},
        }

        applied: list[str] = []
        treatment: dict[str, Any] = {}
        if le_rule:
            applied.append(le_rule.rule_id)
            treatment["fund"] = le_rule.action.get("fund")
            treatment["fund_name"] = le_rule.action.get("fund_name")
        if coa_rule:
            applied.append(coa_rule.rule_id)
            for k in ("account", "account_name", "classification", "trans_type_debit", "batch_type"):
                if coa_rule.action.get(k) is not None:
                    treatment[k] = coa_rule.action[k]
        deal_ids = [r.action.get("deal_id") for r in deal_hits.values() if r]
        for r in deal_hits.values():
            if r:
                applied.append(r.rule_id)
        if deal_ids:
            treatment["deal_ids"] = deal_ids

        complete = bool(le_rule and coa_rule and not missing_deals)

        # Filed snapshot = the verified loader (what the administrator shipped).
        in_upload = g.legal_entity in upload_entities
        agree: Optional[bool] = None
        if coa_rule and in_upload:
            t_default = coa_rule.action.get("classification")
            t_debit = coa_rule.action.get("trans_type_debit") or t_default
            a1 = agreement.get((g.legal_entity, t_debit))
            a2 = agreement.get((g.legal_entity, t_default))
            agree = bool(a1 is not False and a2 is not False)
        filed_snapshot = {
            "in_upload_template": in_upload,
            "loader_trans_types": sorted(
                {k[1] for k in upload_by_le_tt if k[0] == g.legal_entity and k[1] in {treatment.get("classification"), treatment.get("trans_type_debit")}}
            )
            if in_upload
            else [],
            "loader_agrees": agree,
        }

        gap = gaps.get((g.gl_account, g.trans_type))
        source_flags: list[str] = []
        if gap:
            source_flags.append("mapping_gap")
        if g.legal_entity not in entity_listing:
            source_flags.append("entity_miss")
        if any(d not in deals_list for d in g.deals):
            source_flags.append("deal_miss")

        reasons: list[str] = []
        if not le_rule:
            reasons.append("entity_miss")
        if not coa_rule:
            reasons.append("no_matching_rule")
        if missing_deals:
            reasons.append("deal_miss")
        if complete and agree is False:
            reasons.append("rule_mismatch")

        mapping = {k: v for k, v in treatment.items() if v is not None}
        can_be_rule = complete and agree is not False
        mapping_method = None
        rule_id = None
        case_id = None
        status = "unresolved"
        if can_be_rule:
            mapping_method = "rule"
            rule_id = coa_rule.rule_id  # type: ignore[union-attr]
            status = "resolved"
        else:
            case_id = f"CASE-GL-{len(cases) + 1:04d}"
            primary = primary_reason_of(reasons)
            candidates: list[dict[str, Any]] = []
            if gap and (gap.get("proposed_account") or gap.get("proposed_trans_type")):
                candidates.append(
                    {
                        "treatment": f"Mapping Gaps proposal: {gap.get('proposed_account') or gap.get('proposed_trans_type')}",
                        "classification": gap.get("proposed_trans_type"),
                        "account": (gap.get("proposed_account") or "").split(" - ")[0] or None,
                        "account_name": gap.get("proposed_account"),
                        "project_code": None,
                        "counterparty": None,
                        "rationale": f"Proposed on Mapping Gaps row {gap['row']}; approval column blank",
                        "pnl_delta": 0.0,
                        "evidence": [sid[MAPPING_GAPS_SHEET]],
                    }
                )
            if "Partner Transfer" in g.batch_types and not coa_rule:
                candidates.append(
                    {
                        "treatment": "Batch Preference: Partner Transfer batch override",
                        "classification": "Partner Transfer",
                        "account": None,
                        "project_code": None,
                        "counterparty": None,
                        "rationale": "Batch Preference sheet ranks Partner Transfer first (-1); GL rows carry batch type Partner Transfer",
                        "pnl_delta": 0.0,
                        "evidence": [sid["Batch Preference"]] if "Batch Preference" in sid else [],
                    }
                )
            if treatment.get("classification"):
                candidates.append(
                    {
                        "treatment": f"Crosswalk result: {treatment.get('classification')}",
                        "classification": treatment.get("classification"),
                        "account": treatment.get("account"),
                        "project_code": None,
                        "counterparty": None,
                        "rationale": "Deterministic CoA Mapping hit; loader disagrees or entity/deal unmapped",
                        "pnl_delta": 0.0,
                        "evidence": [coa_rule.rule_id] if coa_rule else [],
                    }
                )
            suggestion = (
                {
                    "candidates": candidates[:4],
                    "candidates_note": "Deterministic candidates from Mapping Gaps, Batch Preference and crosswalk evaluation",
                    "drafted_by": "normalize_gl",
                    "model": None,
                }
                if len(candidates) >= 1
                else None
            )
            cases.append(
                Case(
                    case_id=case_id,
                    line_id=line_id,
                    reasons=reasons,
                    primary_reason=primary,
                    status="unresolved",
                    priority=priority_for_reason(primary),
                    source_flags=source_flags,
                    suggestion=suggestion,
                    partial_rule_id=applied[0] if applied else None,
                    recon_category="mapping_gap",
                    materiality_tier=priority_for_reason(primary),
                    clearance_target="Resolve or escalate before loader sign-off",
                    preparer="Fund Accountant",
                    reviewer="Fund Admin",
                )
            )

        excerpt = (
            f"{g.legal_entity} | {g.gl_account} | {g.trans_type} | "
            f"{g.row_count} GL rows | Dr {_money(g.debits):,.2f} Cr {_money(g.credits):,.2f} {g.currency}"
        )
        lines.append(
            Line(
                line_id=line_id,
                close_id=CLOSE_ID,
                date="2026-06-30",
                description=f"{g.gl_account} · {g.trans_type}",
                amount=_money(g.net),
                currency=g.currency,
                source_ref={
                    "source_id": gl_source_id,
                    "location": {"sheet": GL_SHEET, "row": g.first_row, "row_last": g.last_row},
                },
                status=status,
                raw_facts=raw_facts,
                filed_snapshot=filed_snapshot,
                source_flags=source_flags,
                mapping=mapping,
                rule_evaluation={
                    "match_result": "exact_hit" if complete else ("candidates" if applied else "none"),
                    "complete": complete,
                    "rule_id": rule_id,
                    "applied_rule_ids": applied,
                    "treatment": treatment,
                    "blocking_reasons": [r for r in reasons if r != "rule_mismatch"],
                    "comparison": {"agrees": agree, "basis": "loader (LE, trans type) debits/credits"},
                },
                rule_id=rule_id,
                case_id=case_id,
                bank_ref={
                    "source_id": gl_source_id,
                    "location": {"sheet": GL_SHEET, "row": g.first_row},
                    "excerpt": excerpt,
                    "match_method": "exact",
                    "row_count": g.row_count,
                },
                mapping_method=mapping_method,
            )
        )
        evaluation_report.append(
            {
                "line_id": line_id,
                "complete": complete,
                "agrees": agree,
                "reasons": reasons,
                "source_flags": source_flags,
                "mapping_method": mapping_method,
                "row_count": g.row_count,
            }
        )

    # ---- Tie-outs ------------------------------------------------------------
    tieouts: list[TieOut] = []

    def _tie(kind: str, src: Decimal, ln: Decimal, *, note: str, currency: Optional[str] = None, basis: str = "txn", failed: Optional[list] = None, status: Optional[str] = None) -> None:
        diff = src - ln
        st = status or ("passed" if abs(diff) <= TOL and not failed else "failed")
        tieouts.append(
            TieOut(
                tieout_id=f"TIE-GL-{len(tieouts) + 1:03d}",
                close_id=CLOSE_ID,
                source_total=_money(src),
                line_total=_money(ln),
                difference=_money(diff),
                status=st,
                currency=currency,
                kind=kind,
                note=note,
                failed_batches=failed or None,
                amount_basis=basis,
            )
        )

    for ccy in sorted(gl_meta["ccy_totals"]):
        deb, cre = (Decimal(x) for x in gl_meta["ccy_totals"][ccy])
        l_deb, l_cre = line_ccy_totals[ccy]
        _tie(
            "staging_to_lines",
            deb + cre,
            l_deb + l_cre,
            note="GL debits + credits (entity currency) vs canonical Line movement groups",
            currency=ccy,
            basis="abs_rollup",
        )

    bad_batches = [k for k, (deb, cre) in gl_meta["batches"].items() if abs(Decimal(deb) - Decimal(cre)) > TOL]
    gl_deb = sum((Decimal(v[0]) for v in gl_meta["batches"].values()), Decimal("0"))
    gl_cre = sum((Decimal(v[1]) for v in gl_meta["batches"].values()), Decimal("0"))
    _tie(
        "diu_batch_zero_balance",
        gl_deb,
        gl_cre,
        note=f"GL per-batch debit==credit across {len(gl_meta['batches'])} batches (entity currency)",
        failed=bad_batches,
    )

    bad_up = [str(k) for k, (deb, cre) in upload_batches.items() if abs(deb - cre) > TOL]
    _tie(
        "loader_batch_zero_balance",
        sum((v[0] for v in upload_batches.values()), Decimal("0")),
        sum((v[1] for v in upload_batches.values()), Decimal("0")),
        note=f"Verified loader per-batch debit==credit across {len(upload_batches)} batches",
        failed=bad_up,
    )

    net_diffs = [le for le in upload_entities if abs(upload_by_le_net[le] - gl_net_by_le.get(le, Decimal("0"))) > TOL]
    _tie(
        "gl_vs_loader_entity_net",
        sum((gl_net_by_le[le] for le in upload_entities), Decimal("0")),
        sum((upload_by_le_net[le] for le in upload_entities), Decimal("0")),
        note=f"Net movement per entity: GL vs verified loader for {len(upload_entities)} loader entities",
        basis="entity",
        failed=net_diffs,
    )

    residue_actual = {
        "upload_entities_not_in_entity_listing": len([e for e in upload_entities if e not in entity_listing]),
        "upload_deals_not_in_deals_list": len([d for d in upload_deals if d not in deals_list]),
        "investor_mapping_names_not_in_investors_list": len(unmatched_investors),
    }
    for name, expected in EXPECTED_RESIDUE.items():
        actual = residue_actual[name]
        if actual != expected:
            raise ValueError(
                f"INGESTION FAILED: README residue {name} got {actual}, expected {expected}"
            )
        _tie(
            "readme_residue",
            Decimal(expected),
            Decimal(actual),
            note=f"README2 known unmatched rows preserved: {name.replace('_', ' ')}",
            basis="count",
        )
    lines_by_id = {ln.line_id: ln for ln in lines}
    gap_keys_as_cases = {
        (lines_by_id[c.line_id].raw_facts["gl_account"], lines_by_id[c.line_id].raw_facts["trans_type"])
        for c in cases
        if "mapping_gap" in c.source_flags
    }
    _tie(
        "readme_residue",
        Decimal(len(gaps)),
        Decimal(len(gap_keys_as_cases)),
        note="Mapping Gaps sheet rows (account / trans type) each surfaced as at least one Case",
        basis="count",
    )

    close = Close(
        close_id=CLOSE_ID,
        fund_id=FUND_ID,
        period=PERIOD,
        workflow=WORKFLOW,
        status="in_review",
        contract_version=CONTRACT_VERSION,
    )
    model = CanonicalClose(
        close=close,
        sources=sources,
        lines=lines,
        rules=rules,
        cases=cases,
        decisions=[],
        tieouts=tieouts,
    )
    ensure_close_tasks(model)
    model.evaluation_report = evaluation_report  # type: ignore[attr-defined]
    model.gl_summary = {  # type: ignore[attr-defined]
        "gl_rows": gl_meta["row_count"],
        "movement_groups": len(lines),
        "entities": len({g.legal_entity for g in ordered}),
        "loader_entities": len(upload_entities),
        "loader_rows": len(up_rows),
        "unmatched_investor_names": len(unmatched_investors),
        "mapping_gap_rows": len(gaps),
        "residue": residue_actual,
        "reason_counts": dict(Counter(r for c in cases for r in c.reasons)),
    }
    return model
